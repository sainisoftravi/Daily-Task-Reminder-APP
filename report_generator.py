#!/usr/bin/env python3
"""
Daily Task Log Reminder System - Task Fill Report Generator (XLSX & PDF)
Generates user-wise, date-wise matrix reports with dynamic team, period, location headers,
and includes the downloader's full name in the header banner.
"""

import io
import re
import html
import datetime
from typing import List, Dict, Any, Optional

import database
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


def sanitize_str(val: Any) -> str:
    """Removes illegal control characters for Excel/XML compatibility."""
    if val is None:
        return ""
    return re.sub(r'[\x00-\x08\x0B-\x0C\x0E-\x1F]', '', str(val))


def generate_xlsx_report(
    team_name: str,
    period_label: str,
    locations_str: str,
    dates_list: List[str],
    employees_list: List[Dict[str, Any]],
    logs_map: Dict[tuple, str],
    downloaded_by: Optional[str] = None
) -> bytes:
    """
    Generates an Excel (.xlsx) user-wise task fill report matrix.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Task Fill Report"
    ws.views.sheetView[0].showGridLines = True

    total_cols = len(employees_list) + 1
    last_col_letter = get_column_letter(total_cols)

    # 1. Top Title Banner (Merged Row 1)
    banner_text = sanitize_str(f"{team_name} - {period_label} ({locations_str})")
    ws.merge_cells(f"A1:{last_col_letter}1")
    banner_cell = ws["A1"]
    banner_cell.value = banner_text
    banner_cell.font = Font(name="Calibri", size=13, bold=True, color="0F172A")
    banner_cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    banner_cell.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[1].height = 26

    # Sub-header row for Downloader Info
    if downloaded_by:
        ws.merge_cells(f"A2:{last_col_letter}2")
        sub_cell = ws["A2"]
        curr_time_str = datetime.datetime.now().strftime("%d-%b-%Y %I:%M %p")
        sub_cell.value = sanitize_str(f"Downloaded By: {downloaded_by}   |   Generated On: {curr_time_str}")
        sub_cell.font = Font(name="Calibri", size=9.5, italic=True, bold=True, color="475569")
        sub_cell.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        sub_cell.alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[2].height = 20
        header_row_num = 3
    else:
        header_row_num = 2

    # 2. Table Header Row
    ws.cell(row=header_row_num, column=1, value="Date")
    header_fill_date = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid") # Amber
    date_header_cell = ws.cell(row=header_row_num, column=1)
    date_header_cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    date_header_cell.fill = header_fill_date
    date_header_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Distinct header background colors per employee column
    col_palette = [
        "3B82F6", "10B981", "8B5CF6", "EC4899", "06B6D4",
        "6366F1", "F97316", "14B8A6", "64748B", "D946EF"
    ]

    for idx, emp in enumerate(employees_list):
        col_idx = idx + 2
        cell = ws.cell(row=header_row_num, column=col_idx, value=sanitize_str(emp.get("name", "Employee")))
        bg_color = col_palette[idx % len(col_palette)]
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[header_row_num].height = 24

    # Border definitions
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    # Apply border to header row
    for col in range(1, total_cols + 1):
        ws.cell(row=header_row_num, column=col).border = thin_border

    # 3. Data Rows
    start_data_row = header_row_num + 1
    for r_idx, date_str in enumerate(dates_list):
        row_num = r_idx + start_data_row
        ws.row_dimensions[row_num].height = 40

        # Date column
        dt_cell = ws.cell(row=row_num, column=1, value=sanitize_str(date_str))
        dt_cell.font = Font(name="Calibri", size=10, bold=True, color="334155")
        dt_cell.alignment = Alignment(horizontal="center", vertical="top")
        dt_cell.border = thin_border

        # Employee columns
        for c_idx, emp in enumerate(employees_list):
            col_num = c_idx + 2
            emp_name = emp.get("name", "")
            key = (date_str, emp_name.lower())
            details = logs_map.get(key, "")

            cell = ws.cell(row=row_num, column=col_num)
            if details and details.strip():
                det_strip = sanitize_str(details.strip())
                det_lower = det_strip.lower()
                cell.value = det_strip
                if "week off" in det_lower or "weekoff" in det_lower or det_strip.startswith("🏖️"):
                    cell.font = Font(name="Calibri", size=9.5, bold=True, color="0369A1")
                    cell.fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
                elif "on leave" in det_lower or det_strip.startswith("🌴") or "leave" in det_lower:
                    cell.font = Font(name="Calibri", size=9.5, bold=True, color="B45309")
                    cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                else:
                    cell.font = Font(name="Calibri", size=9.5, color="0F172A")
                    cell.fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
            else:
                cell.value = "Data Not Available"
                cell.font = Font(name="Calibri", size=9, italic=True, color="94A3B8")
                cell.fill = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")

            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = thin_border

    # 4. Column Widths
    ws.column_dimensions["A"].width = 14
    for col in range(2, total_cols + 1):
        col_letter = get_column_letter(col)
        ws.column_dimensions[col_letter].width = 30

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


def generate_pdf_report(
    team_name: str,
    period_label: str,
    locations_str: str,
    dates_list: List[str],
    employees_list: List[Dict[str, Any]],
    logs_map: Dict[tuple, str],
    downloaded_by: Optional[str] = None
) -> bytes:
    """
    Generates a PDF document task fill report matrix.
    """
    pdf_buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        pdf_buffer,
        pagesize=landscape(letter),
        leftMargin=24,
        rightMargin=24,
        topMargin=24,
        bottomMargin=24
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=colors.HexColor("#0F172A"),
        alignment=2 # Right aligned
    )
    sub_style = ParagraphStyle(
        'SubStyle',
        parent=styles['Normal'],
        fontName='Helvetica-BoldOblique',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#475569"),
        alignment=2
    )
    header_cell_style = ParagraphStyle(
        'HeaderCellStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=11,
        textColor=colors.white,
        alignment=1
    )
    date_cell_style = ParagraphStyle(
        'DateCellStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=10,
        textColor=colors.HexColor("#334155"),
        alignment=1
    )
    log_text_style = ParagraphStyle(
        'LogTextStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#0F172A")
    )
    weekoff_text_style = ParagraphStyle(
        'WeekOffTextStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#0369A1")
    )
    leave_text_style = ParagraphStyle(
        'LeaveTextStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#B45309")
    )
    pending_text_style = ParagraphStyle(
        'PendingTextStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#94A3B8")
    )

    elements = []

    # 1. Header Banner
    safe_title = html.escape(f"{team_name} - {period_label} ({locations_str})")
    elements.append(Paragraph(f"<b>{safe_title}</b>", title_style))

    if downloaded_by:
        curr_time_str = datetime.datetime.now().strftime("%d-%b-%Y %I:%M %p")
        safe_dl = html.escape(downloaded_by)
        elements.append(Paragraph(f"Downloaded By: {safe_dl} &nbsp;|&nbsp; Generated On: {curr_time_str}", sub_style))

    elements.append(Spacer(1, 8))

    # 2. Build Table Data Matrix & Custom Cell Backgrounds
    header_row = [Paragraph("<b>Date</b>", header_cell_style)]
    for emp in employees_list:
        emp_name_clean = html.escape(emp.get('name', 'Employee'))
        header_row.append(Paragraph(f"<b>{emp_name_clean}</b>", header_cell_style))

    table_data = [header_row]

    # Base table styling
    t_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E293B")),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]

    for r_idx, date_str in enumerate(dates_list):
        row_num = r_idx + 1 # Table row index (1-based after header)
        row = [Paragraph(html.escape(date_str), date_cell_style)]
        
        # Default alternating row background for row
        base_bg = colors.HexColor("#F8FAFC") if r_idx % 2 == 1 else colors.white
        t_style.append(('BACKGROUND', (0, row_num), (0, row_num), base_bg))

        for c_idx, emp in enumerate(employees_list):
            col_num = c_idx + 1 # Table column index (1-based after date)
            emp_name = emp.get("name", "")
            key = (date_str, emp_name.lower())
            details = logs_map.get(key, "")

            if details and details.strip():
                det_strip = details.strip()
                det_lower = det_strip.lower()
                safe_details = html.escape(det_strip).replace("\n", "<br/>")
                
                if "week off" in det_lower or "weekoff" in det_lower or det_strip.startswith("🏖️"):
                    row.append(Paragraph(f"<b>🏖️ {safe_details}</b>", weekoff_text_style))
                    t_style.append(('BACKGROUND', (col_num, row_num), (col_num, row_num), colors.HexColor("#E0F2FE")))
                elif "on leave" in det_lower or det_strip.startswith("🌴") or "leave" in det_lower:
                    row.append(Paragraph(f"<b>🌴 {safe_details}</b>", leave_text_style))
                    t_style.append(('BACKGROUND', (col_num, row_num), (col_num, row_num), colors.HexColor("#FEF3C7")))
                else:
                    row.append(Paragraph(safe_details, log_text_style))
                    t_style.append(('BACKGROUND', (col_num, row_num), (col_num, row_num), colors.HexColor("#F0FDF4")))
            else:
                row.append(Paragraph("Data Not Available", pending_text_style))
                t_style.append(('BACKGROUND', (col_num, row_num), (col_num, row_num), colors.HexColor("#FFFBEB")))
        table_data.append(row)

    printable_width = 744
    date_col_w = 64
    remaining_w = printable_width - date_col_w
    emp_col_w = max(70, remaining_w / max(1, len(employees_list)))
    col_widths = [date_col_w] + [emp_col_w] * len(employees_list)

    pdf_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    pdf_table.setStyle(TableStyle(t_style))
    elements.append(pdf_table)

    doc.build(elements)
    pdf_buffer.seek(0)
    return pdf_buffer.read()
